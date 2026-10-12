"""
Converts the Documentation Agent's Markdown output into a .docx file.
Handles headings, bullet lists, inline **bold**, Markdown tables (rendered
as real Word tables), and skips horizontal-rule lines instead of printing
them literally.
"""

import io
import re
from docx import Document
from docx.shared import Pt

BOLD_PATTERN = re.compile(r"(\*\*.+?\*\*)")

# Matches annotations like "(≈ 150 words)", "(~150 words)", "(150-200 words)".
WORD_COUNT_PATTERN = re.compile(
    r"\s*\(\s*[≈~]?\s*\d+(?:\s*[-–—]\s*\d+)?\s*words?\s*\)", re.IGNORECASE
)


def _add_inline_runs(paragraph, text: str):
    for chunk in BOLD_PATTERN.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        else:
            paragraph.add_run(chunk)


def _tight_paragraph(doc, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    return p


def _is_table_separator(line: str) -> bool:
    return bool(re.match(r"^\|?[\s:|-]+\|?$", line)) and "-" in line


def _parse_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _add_table(doc, rows: list[list[str]]):
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(n_cols):
            text = row[j] if j < len(row) else ""
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            _add_inline_runs(p, text)
            if i == 0:
                for run in p.runs:
                    run.bold = True


def markdown_to_docx_bytes(markdown_text: str, title: str = "Project Documentation") -> bytes:
    doc = Document()

    normal_style = doc.styles["Normal"]
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)

    heading0 = doc.add_heading(title, level=0)
    heading0.paragraph_format.space_after = Pt(10)

    lines = markdown_text.splitlines()
    i = 0
    n = len(lines)

    while i < n:
<<<<<<< HEAD:backend/app/export.py
        line = lines[i].rstrip()
=======
        line = WORD_COUNT_PATTERN.sub("", lines[i]).rstrip()
>>>>>>> 95add18 (Integrated feasibility agent):academic-mentor/backend/app/export.py

        if not line.strip():
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line.strip()):
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < n and _is_table_separator(lines[i + 1].strip()):
            table_rows = [_parse_table_row(line)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                table_rows.append(_parse_table_row(lines[i]))
                i += 1
            _add_table(doc, table_rows)
            continue

        if line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=3)
            h.paragraph_format.space_after = Pt(6)
        elif line.startswith("## "):
            h = doc.add_heading(line[3:].strip(), level=2)
            h.paragraph_format.space_after = Pt(6)
        elif line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.paragraph_format.space_after = Pt(8)
        elif line.strip().startswith(("- ", "* ")):
            p = _tight_paragraph(doc, style="List Bullet")
            _add_inline_runs(p, line.strip()[2:].strip())
        else:
            p = _tight_paragraph(doc)
            _add_inline_runs(p, line.strip())

        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
